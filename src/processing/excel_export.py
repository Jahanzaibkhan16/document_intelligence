# src/processing/excel_export.py
"""Excel export — convert extracted documents to Excel workbooks."""
import json
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

from src.core.config import EXTRACTED_DIR
from src.core.database import list_documents


# ================= STYLE HELPERS =================
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78",
                            fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _style_header_row(ws, row_num, num_cols):
    """Style a header row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _auto_width(ws, max_width=50):
    """Auto-size columns based on content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)


# ================= EXCEL EXPORT =================
def export_to_excel(doc_ids=None, filepath=None):
    """
    Export documents to an Excel workbook.

    Args:
        doc_ids: list of document IDs (None = all documents)
        filepath: output path (None = auto-generate)

    Returns:
        Path to saved Excel file
    """
    # Fetch documents
    if doc_ids:
        from src.core.database import get_document
        docs = [get_document(did) for did in doc_ids]
        docs = [d for d in docs if d]
    else:
        docs = list_documents(limit=10000)

    if not docs:
        raise ValueError("No documents to export")

    # Create workbook
    wb = Workbook()

    # ============ SHEET 1 — SUMMARY ============
    ws = wb.active
    ws.title = "Summary"

    # Title
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "Document Extraction Report"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(italic=True, size=10, color="666666")

    # Headers
    headers = [
        "ID", "Filename", "Type", "Status",
        "Confidence", "Uploaded", "Key Fields", "Notes"
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=4, column=col, value=h)
    _style_header_row(ws, 4, len(headers))

    # Data rows
    for i, doc in enumerate(docs, 1):
        row = 4 + i

        fields = {}
        if doc.get('extracted_json'):
            try:
                fields = json.loads(doc['extracted_json'])
            except Exception:
                pass

        # Key fields summary (top 3 non-null)
        key_fields = []
        for k in ['invoice_number', 'vendor_name', 'total_amount',
                  'store_name', 'id_number', 'full_name']:
            if fields.get(k):
                key_fields.append(f"{k.replace('_', ' ').title()}: {fields[k]}")
                if len(key_fields) >= 3:
                    break

        ws.cell(row=row, column=1, value=doc['id']).border = BORDER
        ws.cell(row=row, column=2, value=doc['filename']).border = BORDER
        ws.cell(row=row, column=3, value=doc['doc_type'] or '-').border = BORDER
        ws.cell(row=row, column=4, value=doc['status']).border = BORDER
        ws.cell(row=row, column=5,
                value=f"{(doc['confidence'] or 0)*100:.0f}%"
                ).border = BORDER
        ws.cell(row=row, column=6,
                value=str(doc['uploaded_on'])[:16]).border = BORDER
        ws.cell(row=row, column=7,
                value=" | ".join(key_fields)[:100]).border = BORDER
        ws.cell(row=row, column=8,
                value=doc.get('error_message') or '').border = BORDER

    _auto_width(ws)

    # ============ SHEET 2 — LINE ITEMS ============
    ws2 = wb.create_sheet("Line Items")

    headers2 = [
        "Doc ID", "Doc Type", "Invoice/Receipt #",
        "Item Description", "Quantity", "Unit Price", "Total"
    ]
    for col, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=col, value=h)
    _style_header_row(ws2, 1, len(headers2))

    row2 = 2
    for doc in docs:
        if not doc.get('extracted_json'):
            continue

        try:
            fields = json.loads(doc['extracted_json'])
        except Exception:
            continue

        # Find line items (both invoice + receipt use different keys)
        items = fields.get('line_items') or fields.get('items') or []
        if not items:
            continue

        doc_num = (fields.get('invoice_number') or
                    fields.get('receipt_number') or
                    fields.get('store_name') or
                    doc['filename'])

        for item in items:
            if not isinstance(item, dict):
                continue

            ws2.cell(row=row2, column=1, value=doc['id']).border = BORDER
            ws2.cell(row=row2, column=2,
                     value=doc['doc_type'] or '-').border = BORDER
            ws2.cell(row=row2, column=3, value=str(doc_num)).border = BORDER
            ws2.cell(row=row2, column=4,
                     value=item.get('description') or item.get('name') or ''
                     ).border = BORDER
            ws2.cell(row=row2, column=5,
                     value=item.get('quantity') or '').border = BORDER
            ws2.cell(row=row2, column=6,
                     value=item.get('unit_price') or item.get('price') or ''
                     ).border = BORDER
            ws2.cell(row=row2, column=7,
                     value=item.get('total') or '').border = BORDER
            row2 += 1

    if row2 == 2:
        ws2.cell(row=2, column=1,
                 value="No line items found in any document.")

    _auto_width(ws2)

    # ============ SHEET 3 — BY TYPE ============
    ws3 = wb.create_sheet("By Type")

    type_counts = {}
    type_amounts = {}

    for doc in docs:
        t = doc['doc_type'] or 'unknown'
        type_counts[t] = type_counts.get(t, 0) + 1

        if doc.get('extracted_json'):
            try:
                fields = json.loads(doc['extracted_json'])
                amt = fields.get('total_amount') or fields.get('total') or 0
                try:
                    amt = float(amt)
                    type_amounts[t] = type_amounts.get(t, 0) + amt
                except (ValueError, TypeError):
                    pass
            except Exception:
                pass

    headers3 = ["Document Type", "Count", "Total Amount"]
    for col, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=col, value=h)
    _style_header_row(ws3, 1, len(headers3))

    row3 = 2
    for t in sorted(type_counts.keys()):
        ws3.cell(row=row3, column=1,
                 value=t.replace('_', ' ').title()).border = BORDER
        ws3.cell(row=row3, column=2, value=type_counts[t]).border = BORDER
        ws3.cell(row=row3, column=3,
                 value=round(type_amounts.get(t, 0), 2)).border = BORDER
        row3 += 1

    # Totals row
    ws3.cell(row=row3, column=1, value="TOTAL").font = Font(bold=True)
    ws3.cell(row=row3, column=2, value=sum(type_counts.values())
             ).font = Font(bold=True)
    ws3.cell(row=row3, column=3,
             value=round(sum(type_amounts.values()), 2)
             ).font = Font(bold=True)

    _auto_width(ws3)

    # ============ SAVE ============
    if filepath is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = EXTRACTED_DIR / f"export_{timestamp}.xlsx"

    wb.save(filepath)
    return str(filepath)


# ================= TEST =================
if __name__ == "__main__":
    print("=" * 60)
    print("EXCEL EXPORT TEST")
    print("=" * 60)

    try:
        path = export_to_excel()
        print(f"✅ Exported to: {path}")
        print(f"   Open with Excel or LibreOffice")
    except ValueError as e:
        print(f"⚠️ {e}")
        print("   Upload some documents first via the Streamlit app")
    except Exception as e:
        print(f"❌ Error: {e}")